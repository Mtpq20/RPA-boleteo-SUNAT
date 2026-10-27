"""Lectura/escritura del Excel. Una fila de 'Ventas' = un producto vendido.
Las filas con la misma Ref_Odoo forman UNA boleta (una orden con varios productos);
una fila sin Ref_Odoo es una boleta de un solo producto."""
import hashlib
from copy import copy
from datetime import datetime
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

COLUMNAS = ["ID", "Fecha", "Cuenta", "Tipo_Doc", "Nro_Doc", "Cliente", "Descripcion",
            "Cantidad", "Precio_Unit", "Total", "Estado", "Serie_Numero", "Fecha_Emision", "Error", "Ref_Odoo"]
# Tipo_Doc: 1=DNI, 0=Sin documento (solo si total < S/700), 6=RUC
#   Si Tipo_Doc está vacío se deduce de Nro_Doc: vacío -> 0 (solo nombre), 8 dígitos -> DNI, 11 -> RUC.
# Precio_Unit incluye IGV.
CAMPOS_VENTA = ["Cuenta", "Tipo_Doc", "Nro_Doc", "Cliente", "Descripcion", "Cantidad", "Precio_Unit"]


def normalizar_doc(d):
    nd = d["Nro_Doc"]
    # Excel guarda el DNI como número y se come los ceros: 01234567 -> 1234567
    nd = "" if nd is None else str(int(nd)) if isinstance(nd, float) else str(nd).strip()
    if set(nd) == {"0"}:   # Odoo pone 00000000 a "Consumidor Final": es sin documento
        nd = ""
    td = "" if d["Tipo_Doc"] is None else str(d["Tipo_Doc"]).strip().removesuffix(".0")
    if td == "":
        td = "0" if nd == "" else "6" if len(nd) == 11 else "1"
    if td == "1" and nd.isdigit() and len(nd) < 8:
        nd = nd.zfill(8)
    d["Tipo_Doc"], d["Nro_Doc"] = td, nd


def agregar_lista_cuentas(ws, cuentas):
    """Lista desplegable en la columna Cuenta: evita escribir otra cosa (p. ej. la fecha)."""
    dv = DataValidation(type="list", formula1='"' + ",".join(cuentas) + '"', allow_blank=True,
                        showErrorMessage=True, errorTitle="Cuenta no válida",
                        error="Elige una cuenta de la lista: " + ", ".join(cuentas))
    ws.add_data_validation(dv)
    dv.add("C2:C5000")


def crear_plantilla(ruta, cuentas, filas_demo=()):
    wb = Workbook()
    ws = wb.active
    ws.title = "Ventas"
    ws.append(COLUMNAS)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1F4E78")
    for i, f in enumerate(filas_demo, start=2):
        ws.append(list(f) + [f"=ROUND(H{i}*I{i},2)", "PENDIENTE", None, None, None, None])
    for i in range(1, len(COLUMNAS) + 1):
        ws.column_dimensions[get_column_letter(i)].width = 16
    ws.column_dimensions["G"].width = 32
    ws.freeze_panes = "A2"
    agregar_lista_cuentas(ws, cuentas)

    escribir_resumen(wb.create_sheet("Resumen"), cuentas)
    wb.save(ruta)


def escribir_resumen(rs, cuentas):
    """Hoja Resumen por cuenta. Una boleta puede ocupar varias filas (una por producto):
    se cuentan números de boleta distintos, no filas."""
    rs.delete_rows(1, rs.max_row)
    rs.append(["Cuenta", "N° boletas emitidas", "Total boleteado (S/)", "Productos pendientes",
               "Productos con error", "Total pendiente (S/)"])
    for c in rs[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1F4E78")
    rango = lambda col: f"Ventas!${col}$2:${col}$5000"
    for n, cuenta in enumerate(cuentas, start=2):
        rs.append([
            cuenta,
            f'=SUMPRODUCT(({rango("C")}=A{n})*({rango("K")}="EMITIDA")'
            f'/COUNTIFS({rango("L")},{rango("L")}&"",{rango("C")},{rango("C")}&""))',
            f'=SUMIFS(Ventas!$J:$J,Ventas!$C:$C,A{n},Ventas!$K:$K,"EMITIDA")',
            f'=COUNTIFS(Ventas!$C:$C,A{n},Ventas!$K:$K,"PENDIENTE")',
            f'=COUNTIFS(Ventas!$C:$C,A{n},Ventas!$K:$K,"ERROR")',
            f'=SUMIFS(Ventas!$J:$J,Ventas!$C:$C,A{n},Ventas!$K:$K,"PENDIENTE")',
        ])
    t = len(cuentas) + 2
    rs.append(["TOTAL"] + [f"=SUM({get_column_letter(c)}2:{get_column_letter(c)}{t-1})" for c in range(2, 7)])
    for c in rs[t]:
        c.font = Font(bold=True)
    for r in rs.iter_rows(min_row=2, min_col=3, max_col=3):
        r[0].number_format = '#,##0.00'
    for r in rs.iter_rows(min_row=2, min_col=6, max_col=6):
        r[0].number_format = '#,##0.00'
    for i in range(1, 7):
        rs.column_dimensions[get_column_letter(i)].width = 24


def asegurar_columna(ws, enc, nombre):
    """Agrega al final una columna que falte en un Excel creado con una versión anterior."""
    if nombre not in enc:
        col = len(enc) + 1
        ws.cell(1, col, nombre)
        ws.cell(1, col).font = copy(ws.cell(1, col - 1).font)
        ws.cell(1, col).fill = copy(ws.cell(1, col - 1).fill)
        enc.append(nombre)


def huella(ws, idx, fila):
    """Resumen de los datos de la venta: si cambia, la fila en ERROR se vuelve a intentar."""
    datos = "|".join(str(ws.cell(fila, idx[c] + 1).value) for c in CAMPOS_VENTA)
    return hashlib.md5(datos.encode("utf-8")).hexdigest()[:12]


def leer_pendientes(ruta):
    """Devuelve (wb, ws, idx, pendientes): filas PENDIENTE/vacías, y filas ERROR cuyos datos
    cambiaron desde el error (ERROR = nunca se llegó a emitir, reintentar es seguro).
    REVISAR nunca se reintenta solo."""
    wb = load_workbook(ruta)
    ws = wb["Ventas"]
    enc = [c.value for c in ws[1]]
    asegurar_columna(ws, enc, "Ref_Odoo")
    asegurar_columna(ws, enc, "Huella")   # columna técnica para detectar filas corregidas
    idx = {n: i for i, n in enumerate(enc)}
    pend = []
    for r in range(2, ws.max_row + 1):
        estado = ws.cell(r, idx["Estado"] + 1).value
        if ws.cell(r, 1).value is None:
            continue
        corregida = estado == "ERROR" and ws.cell(r, idx["Huella"] + 1).value != huella(ws, idx, r)
        if estado in (None, "", "PENDIENTE") or corregida:
            d = {n: ws.cell(r, i + 1).value for n, i in idx.items()}
            d["fila"] = r
            # El Total es fórmula; si no hay valor cacheado lo calculamos
            d["Total"] = round(float(d["Cantidad"]) * float(d["Precio_Unit"]), 2)
            pend.append(d)
    # Si se corrigió una fila de una orden en ERROR, se reintenta la orden completa
    refs = {str(d.get("Ref_Odoo") or "").strip() for d in pend} - {""}
    ya = {d["fila"] for d in pend}
    for r in range(2, ws.max_row + 1):
        if (r not in ya and ws.cell(r, idx["Estado"] + 1).value == "ERROR"
                and str(ws.cell(r, idx["Ref_Odoo"] + 1).value or "").strip() in refs):
            d = {n: ws.cell(r, i + 1).value for n, i in idx.items()}
            d["fila"] = r
            d["Total"] = round(float(d["Cantidad"]) * float(d["Precio_Unit"]), 2)
            pend.append(d)
    pend.sort(key=lambda d: d["fila"])
    return wb, ws, idx, pend


def agrupar(ws, idx, pend):
    """Junta las filas pendientes en boletas: misma Ref_Odoo = una boleta con varios productos.
    Cuenta/Cliente/DNI se toman de la fila de la orden que los tenga (basta con la primera)."""
    grupos = {}
    for d in pend:
        ref = str(d.get("Ref_Odoo") or "").strip()
        grupos.setdefault(ref or f"fila {d['fila']}", []).append(d)
    boletas = []
    for clave, filas in grupos.items():
        b = {"clave": clave, "fila": filas[0]["fila"], "filas": [d["fila"] for d in filas], "error": None}
        for campo in ("Cuenta", "Tipo_Doc", "Nro_Doc", "Cliente"):
            # texto -> valor original (el DNI puede venir como número de Excel)
            valores = {str(d[campo]).strip(): d[campo] for d in filas if str(d[campo] or "").strip()}
            b[campo] = next(iter(valores.values())) if len(valores) == 1 else None
            if len(valores) > 1:
                b["error"] = f"Las filas de {clave} tienen distinto {campo}: {sorted(valores)}"
        b["items"] = [{k: d[k] for k in ("Descripcion", "Cantidad", "Precio_Unit", "Total")} for d in filas]
        b["Total"] = round(sum(i["Total"] for i in b["items"]), 2)
        if not clave.startswith("fila "):
            # Toda la orden debe estar pendiente: si una parte ya se emitió, no emitir el resto solo
            en_hoja = [r for r in range(2, ws.max_row + 1) if str(ws.cell(r, idx["Ref_Odoo"] + 1).value or "").strip() == clave]
            if sorted(en_hoja) != sorted(b["filas"]):
                b["error"] = f"La orden {clave} tiene filas en otro estado (filas {en_hoja}); revísala"
        normalizar_doc(b)
        boletas.append(b)
    return boletas


def marcar(wb, ws, idx, ruta, fila, estado, serie_numero=None, error=None):
    # .value = ... (ws.cell(f, c, None) no borra el valor anterior)
    ws.cell(fila, idx["Estado"] + 1).value = estado
    ws.cell(fila, idx["Serie_Numero"] + 1).value = serie_numero
    ws.cell(fila, idx["Fecha_Emision"] + 1).value = datetime.now().strftime("%Y-%m-%d %H:%M:%S") if estado == "EMITIDA" else None
    ws.cell(fila, idx["Error"] + 1).value = error
    ws.cell(fila, idx["Huella"] + 1).value = huella(ws, idx, fila) if estado == "ERROR" else None
    wb.save(ruta)   # se guarda en cada boleta: si el bot se cae, no se pierde nada
