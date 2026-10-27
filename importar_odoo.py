"""Importa a ventas.xlsx las órdenes exportadas desde Odoo (Punto de venta > Órdenes > Exportar).

Uso:  python importar_odoo.py "Órdenes del punto de venta (pos.order).xlsx" [--prueba]
      --prueba  muestra lo que se importaría, sin tocar ventas.xlsx

La exportación debe traer (plantilla "Boletas" en Odoo): Ref. de la orden, Fecha, Punto de venta,
Cliente, Cliente/Número de Identificación, Total y, de Líneas de la orden: Producto, Cantidad,
Precio unitario e Impuestos incl.

Cada producto queda en una fila PENDIENTE con su Ref_Odoo; main.py junta las filas de una misma
orden en UNA boleta. Una orden que ya está en ventas.xlsx no se vuelve a importar.
"""
import os
import sys
import unicodedata
from pathlib import Path

import yaml
from openpyxl import load_workbook

from excel_io import asegurar_columna, escribir_resumen

BASE = Path(__file__).parent


def clave(texto):
    t = unicodedata.normalize("NFKD", str(texto or "")).encode("ascii", "ignore").decode()
    return " ".join(t.lower().split())


# Columna de la exportación de Odoo -> textos que puede tener su encabezado
COLUMNAS_ODOO = {
    "ref": ["ref. de la orden"],
    "fecha": ["fecha"],
    "pdv": ["punto de venta"],
    "cliente": ["cliente"],
    "doc": ["cliente/numero de identificacion", "cliente/nif", "cliente/numero de documento"],
    "total": ["total"],
    "producto": ["lineas de la orden/producto"],
    "cantidad": ["lineas de la orden/cantidad"],
    "precio": ["lineas de la orden/precio unitario"],
    "subtotal": ["lineas de la orden/impuestos incl.", "lineas de la orden/subtotal impuestos incluidos",
                 "lineas de la orden/subtotal con impuestos"],
}
OBLIGATORIAS = ["ref", "pdv", "producto", "cantidad", "precio"]


def leer_exportacion(ruta):
    """Devuelve {ref: {datos de la orden, "lineas": [...]}} en el orden del archivo."""
    ws = load_workbook(ruta, data_only=True).active
    enc = [clave(c.value) for c in ws[1]]
    col = {}
    for campo, nombres in COLUMNAS_ODOO.items():
        col[campo] = next((enc.index(n) for n in nombres if n in enc), None)
    faltan = [c for c in OBLIGATORIAS if col[c] is None]
    if faltan:
        sys.exit(f"A la exportación le faltan columnas: {', '.join(COLUMNAS_ODOO[c][0] for c in faltan)}")

    ordenes, actual = {}, None
    for fila in ws.iter_rows(min_row=2, values_only=True):
        v = {c: (fila[i] if i is not None else None) for c, i in col.items()}
        if v["ref"]:   # primera línea de una orden: trae los datos generales
            actual = ordenes.setdefault(str(v["ref"]).strip(), {**v, "lineas": []})
        if actual is not None and v["producto"]:
            actual["lineas"].append(v)
    return ordenes


def main():
    args = [a for a in sys.argv[1:] if not a.startswith("--")]
    prueba = "--prueba" in sys.argv
    if not args:
        sys.exit(__doc__)
    cfg = yaml.safe_load(open(BASE / "config.yaml", encoding="utf-8"))
    pdv_cuenta = (cfg.get("odoo") or {}).get("puntos_de_venta") or {}
    ruta_ventas = BASE / os.environ.get("EXCEL", cfg["excel"])

    ordenes = leer_exportacion(args[0])
    wb = load_workbook(ruta_ventas)
    ws = wb["Ventas"]
    enc = [c.value for c in ws[1]]
    asegurar_columna(ws, enc, "Ref_Odoo")
    asegurar_columna(ws, enc, "Huella")
    idx = {n: i + 1 for i, n in enumerate(enc)}

    ultima = max((r for r in range(2, ws.max_row + 1) if ws.cell(r, 1).value is not None), default=1)
    ids = [ws.cell(r, 1).value for r in range(2, ultima + 1)]
    sig_id = max([i for i in ids if isinstance(i, (int, float))], default=0) + 1
    ya = {str(ws.cell(r, idx["Ref_Odoo"]).value).strip() for r in range(2, ultima + 1) if ws.cell(r, idx["Ref_Odoo"]).value}

    fila = ultima + 1
    importadas = 0
    for ref, o in ordenes.items():
        if ref in ya:
            print(f"  --   {ref}: ya estaba en {ruta_ventas.name}, no se importa otra vez")
            continue
        cuenta = pdv_cuenta.get(str(o["pdv"] or "").strip())
        if not cuenta:
            print(f"  ERR  {ref}: el punto de venta {o['pdv']!r} no está en config.yaml (odoo > puntos_de_venta)")
            continue
        lineas, avisos = [], []
        for l in o["lineas"]:
            cant = float(l["cantidad"] or 0)
            subtotal = float(l["subtotal"]) if l["subtotal"] is not None else cant * float(l["precio"] or 0)
            if cant <= 0 or subtotal <= 0:
                avisos.append(f"se omitió {l['producto']!r} (cantidad o importe 0: regalo o devolución)")
                continue
            # Precio con IGV y con descuento ya aplicado = lo que pagó el cliente por unidad
            precio = round(subtotal / cant, 2)
            if abs(precio * cant - subtotal) > 0.005:
                avisos.append(f"{l['producto']!r}: {cant:g} x {precio} no da exacto {subtotal} (descuento con céntimos)")
            lineas.append((l["producto"], cant, precio))
        if not lineas:
            print(f"  ERR  {ref}: no tiene productos con importe; no se importa")
            continue
        suma = round(sum(c * p for _, c, p in lineas), 2)
        if o["total"] is not None and abs(suma - float(o["total"])) > 0.01:
            avisos.append(f"la suma de productos S/ {suma:.2f} no coincide con el total de Odoo S/ {float(o['total']):.2f}")

        doc = str(o["doc"] or "").strip()
        doc = None if not doc or set(doc) == {"0"} else doc   # 00000000 = Consumidor Final
        cliente = str(o["cliente"] or "").strip() or None
        print(f"  OK   {ref} -> {cuenta} | {cliente or 'sin nombre'} | {'DNI/RUC ' + doc if doc else 'sin documento'} "
              f"| {len(lineas)} producto(s) | S/ {suma:.2f}")
        for a in avisos:
            print(f"       ! {a}")
        for n, (producto, cant, precio) in enumerate(lineas):
            valores = {
                "ID": sig_id, "Fecha": o["fecha"], "Cuenta": cuenta,
                # Cliente y documento solo en la primera fila: es una sola boleta
                "Nro_Doc": doc if n == 0 else None, "Cliente": cliente if n == 0 else None,
                "Descripcion": str(producto).strip(), "Cantidad": cant, "Precio_Unit": precio,
                "Total": f"=ROUND(H{fila}*I{fila},2)", "Estado": "PENDIENTE", "Ref_Odoo": ref,
            }
            for campo, valor in valores.items():
                ws.cell(fila, idx[campo]).value = valor
            sig_id += 1
            fila += 1
        importadas += 1

    if prueba:
        print(f"\n(--prueba) Se importarían {importadas} orden(es). No se modificó {ruta_ventas.name}.")
        return
    escribir_resumen(wb["Resumen"], list(cfg["cuentas"]))
    try:
        wb.save(ruta_ventas)
    except PermissionError:
        sys.exit(f"No pude guardar {ruta_ventas.name}: ciérralo en Excel y vuelve a ejecutar.")
    print(f"\n{importadas} orden(es) importada(s) a {ruta_ventas.name} como PENDIENTE. "
          f"Revisa cliente/DNI y luego ejecuta: python main.py")


if __name__ == "__main__":
    main()
